"""Iter 337 — Overtime tier + hours-worked timesheets + XLSX template.

Verifies:
  1. `_compute_base_gross` applies OT above the weekly threshold with the
     configured multiplier. Threshold scales with the period.
  2. Timesheet submit accepts `hours_worked` and stores it on the doc.
  3. Preview honours the timesheet's hours_worked override for hourly staff.
  4. XLSX template endpoint returns a valid .xlsx binary with the expected
     header row.
  5. XLSX upload matches by badge_number and creates timesheet drafts.
"""
import os
import sys
import uuid
import asyncio
import io
from pathlib import Path

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env")
sys.path.insert(0, "/app/backend")


def test_iter337_overtime_biweekly():
    """Hourly rate 10,000; OT threshold 40h/wk → biweekly threshold 80h.
    Working 90h: 80h regular + 10h × 1.5 OT.
    Gross = 10,000*80 + 10,000*1.5*10 = 800,000 + 150,000 = 950,000."""
    from routers.hr import _compute_base_gross
    sal = {
        "wage_type": "hourly", "hourly_rate": 10_000, "base_salary": 0,
        "pay_frequency": "biweekly",
        "ot_threshold_hours": 40, "ot_multiplier": 1.5,
    }
    period = "2026-09-16_2026-09-29 (W38-W39)"
    gross, details = _compute_base_gross(sal, period, working_days=10, days_worked=None, hours_worked=90)
    assert gross == 950_000.0, f"expected 950000, got {gross}"
    assert "reg" in details and "OT" in details


def test_iter337_no_overtime_when_disabled():
    from routers.hr import _compute_base_gross
    sal = {"wage_type": "hourly", "hourly_rate": 10_000, "base_salary": 0,
           "pay_frequency": "biweekly"}
    period = "2026-09-16_2026-09-29 (W38-W39)"
    gross, _ = _compute_base_gross(sal, period, working_days=10, days_worked=None, hours_worked=90)
    # No threshold set → flat rate × hours
    assert gross == 900_000.0


def test_iter337_no_overtime_when_under_threshold():
    from routers.hr import _compute_base_gross
    sal = {"wage_type": "hourly", "hourly_rate": 10_000, "base_salary": 0,
           "pay_frequency": "biweekly",
           "ot_threshold_hours": 40, "ot_multiplier": 1.5}
    period = "2026-09-16_2026-09-29 (W38-W39)"
    gross, _ = _compute_base_gross(sal, period, working_days=10, days_worked=None, hours_worked=80)
    assert gross == 800_000.0


def test_iter337_e2e_flow():
    async def _t():
        import deps
        from motor.motor_asyncio import AsyncIOMotorClient
        client = AsyncIOMotorClient(os.environ["MONGO_URL"])
        deps.db = client[os.environ["DB_NAME"]]
        db = deps.db
        from routers.hr import submit_timesheet, preview_payslips
        from routers.hr_timesheet_templates import timesheet_template, timesheet_upload

        loc_id = f"loc_test_{uuid.uuid4().hex[:8]}"
        staff_id = f"u_test_{uuid.uuid4().hex[:8]}"
        sal_id = f"sal_{uuid.uuid4().hex[:8]}"
        badge = f"B{uuid.uuid4().hex[:6].upper()}"
        await db.users.insert_one({
            "id": staff_id, "name": "Hourly Test", "role": "Staff",
            "location_id": loc_id, "email": f"{staff_id}@test.local",
            "badge_number": badge,
        })
        await db.hr_salaries.insert_one({
            "id": sal_id, "staff_id": staff_id, "staff_name": "Hourly Test",
            "base_salary": 0, "hourly_rate": 10_000,
            "currency": "UGX", "pay_frequency": "biweekly",
            "wage_type": "hourly", "location_id": loc_id,
            "ot_threshold_hours": 40, "ot_multiplier": 1.5,
            "line_items": [], "status": "active",
        })
        try:
            # (2) Submit a timesheet with hours_worked
            director = {"id": "u_admin", "name": "Admin", "role": "admin"}
            period = "2026-09-16_2026-09-29 (W38-W39)"
            ts = await submit_timesheet(
                {"period": period, "staff_id": staff_id, "days_worked": 10, "hours_worked": 90},
                director,
            )
            assert ts["hours_worked"] == 90.0
            # Approve so preview picks it up
            await db.hr_timesheets.update_one({"id": ts["id"]}, {"$set": {"status": "approved"}})

            # (3) Preview should use hours_worked and apply OT
            res = await preview_payslips(
                {"period": period, "location_id": loc_id, "use_timesheets": True},
                director,
            )
            row = next(r for r in res["rows"] if r["staff_id"] == staff_id)
            assert row["gross"] == 950_000.0, f"gross={row['gross']}"

            # (4) XLSX template
            resp = await timesheet_template(period=period, location_id=loc_id, current_user=director)
            body = b""
            async for chunk in resp.body_iterator:
                body += chunk if isinstance(chunk, bytes) else chunk.encode()
            assert body.startswith(b"PK"), "should be a valid XLSX (zip magic)"
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(body))
            ws = wb.active
            # Row 5 has the header; confirm our staff shows up beneath
            headers = [ws.cell(row=5, column=c).value for c in range(1, 10)]
            assert "Staff Name" in headers and "Badge Number" in headers
            # Find our staff by badge in column 2
            found = False
            for r in range(6, ws.max_row + 1):
                if str(ws.cell(row=r, column=2).value or "") == badge:
                    found = True
                    break
            assert found, "seeded staff badge should appear in the template"

            # (5) XLSX upload — build a mini-workbook with our staff row filled
            from openpyxl import Workbook
            wb2 = Workbook()
            ws2 = wb2.active
            hdr = ["Staff Name", "Badge Number", "Wage Type", "Rate",
                   "Days Worked", "Hours Worked", "PTO Days", "Notes", "Signature"]
            for c, h in enumerate(hdr, start=1):
                ws2.cell(row=1, column=c, value=h)
            data = ["Hourly Test", badge, "Hourly", "10,000", 10, 100, 0, "OT week", "signed"]
            for c, v in enumerate(data, start=1):
                ws2.cell(row=2, column=c, value=v)
            buf = io.BytesIO(); wb2.save(buf); buf.seek(0)
            class FakeUploadFile:
                filename = "sheet.xlsx"
                async def read(self): return buf.getvalue()
            up = await timesheet_upload(
                file=FakeUploadFile(), period=period, location_id=loc_id, current_user=director,
            )
            assert up["created"], f"expected upload to create a row, got {up}"
            assert up["created"][0]["hours_worked"] == 100.0
            # New draft should now exist
            latest = await db.hr_timesheets.find({"staff_id": staff_id, "period": period}).to_list(10)
            assert any(t.get("hours_worked") == 100.0 for t in latest)
        finally:
            await db.users.delete_one({"id": staff_id})
            await db.hr_salaries.delete_one({"id": sal_id})
            await db.hr_timesheets.delete_many({"staff_id": staff_id})
    asyncio.run(_t())
