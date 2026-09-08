"""Iter 338 — Kiosk autofill + per-day Daily Log XLSX.

Verifies:
  1. `sync_kiosk_to_timesheet(staff_id, 'checkin', ts_iso)` creates a
     timesheet with an open entry when none exists.
  2. Follow-up `checkout` closes the entry and rolls up hours + days.
  3. Non-payroll members are silently ignored (return None).
  4. Downloaded XLSX contains a "Daily Log" sheet with per-day rows.
  5. Uploaded XLSX with Daily Log hours flows into the timesheet's
     hours_worked field even when Summary is blank.
"""
import os
import sys
import io
import uuid
import asyncio
from pathlib import Path

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env")
sys.path.insert(0, "/app/backend")


def test_iter338_kiosk_autofill_and_daily_log():
    async def _t():
        import deps
        from motor.motor_asyncio import AsyncIOMotorClient
        from datetime import datetime, timezone, timedelta
        client = AsyncIOMotorClient(os.environ["MONGO_URL"])
        deps.db = client[os.environ["DB_NAME"]]
        db = deps.db
        from routers.hr import sync_kiosk_to_timesheet
        from routers.hr_timesheet_templates import timesheet_template, timesheet_upload

        loc_id = f"loc_test_{uuid.uuid4().hex[:8]}"
        staff_id = f"u_test_{uuid.uuid4().hex[:8]}"
        sal_id = f"sal_{uuid.uuid4().hex[:8]}"
        badge = f"B{uuid.uuid4().hex[:6].upper()}"
        random_id = f"u_test_{uuid.uuid4().hex[:8]}"

        await db.users.insert_many([
            {"id": staff_id, "name": "Kiosk Test", "role": "Staff",
             "location_id": loc_id, "email": f"{staff_id}@t.local",
             "badge_number": badge},
            {"id": random_id, "name": "Visitor Not On Payroll", "role": "Guest",
             "location_id": loc_id, "email": f"{random_id}@t.local"},
        ])
        await db.hr_salaries.insert_one({
            "id": sal_id, "staff_id": staff_id, "staff_name": "Kiosk Test",
            "base_salary": 0, "hourly_rate": 10_000, "wage_type": "hourly",
            "currency": "UGX", "pay_frequency": "monthly",
            "location_id": loc_id, "line_items": [], "status": "active",
        })
        # HR settings — monthly so the period label is YYYY-MM
        await db.hr_settings.replace_one(
            {"location_id": loc_id},
            {"location_id": loc_id, "hr_enabled": True, "pay_frequency": "monthly", "pay_day": 28},
            upsert=True,
        )
        try:
            now = datetime.now(timezone.utc)
            checkin_iso = (now - timedelta(hours=8)).isoformat()
            checkout_iso = now.isoformat()

            # (1) Check-in creates timesheet with open entry
            ts = await sync_kiosk_to_timesheet(staff_id, "checkin", checkin_iso, loc_id)
            assert ts is not None
            assert ts["period"] == now.strftime("%Y-%m")
            assert len(ts["entries"]) == 1
            assert ts["entries"][0]["check_out_time"] is None
            assert ts.get("hours_worked", 0) == 0  # entry still open

            # (2) Check-out closes it and rolls up hours
            ts2 = await sync_kiosk_to_timesheet(staff_id, "checkout", checkout_iso, loc_id)
            assert ts2["id"] == ts["id"], "same timesheet doc"
            assert ts2["entries"][0]["check_out_time"] == checkout_iso
            # ~8 hours logged (allow tiny drift)
            assert 7.9 <= ts2["hours_worked"] <= 8.1, f"hours={ts2['hours_worked']}"
            assert ts2["days_worked"] == 1

            # (3) Non-payroll member returns None (silent no-op)
            skip = await sync_kiosk_to_timesheet(random_id, "checkin", checkin_iso, loc_id)
            assert skip is None

            # (4) XLSX has a "Daily Log" sheet
            director = {"id": "u_admin", "name": "Admin", "role": "admin"}
            resp = await timesheet_template(period=now.strftime("%Y-%m"), location_id=loc_id, current_user=director)
            body = b""
            async for chunk in resp.body_iterator:
                body += chunk if isinstance(chunk, bytes) else chunk.encode()
            from openpyxl import load_workbook, Workbook
            wb = load_workbook(io.BytesIO(body))
            assert "Daily Log" in wb.sheetnames
            ws_d = wb["Daily Log"]
            # Header on row 3
            headers = [ws_d.cell(row=3, column=c).value for c in range(1, 8)]
            assert "Arrival" in headers and "Exit" in headers and "Hours" in headers

            # (5) Upload — build a workbook where Summary hours is blank
            # but Daily Log has 2 filled days (08:00-17:00 = 9h each = 18h).
            wb2 = Workbook()
            ws2 = wb2.active
            ws2.title = "Timesheet"
            hdr = ["Staff Name", "Badge Number", "Wage Type", "Rate",
                   "Days Worked", "Hours Worked", "PTO Days", "Notes", "Signature"]
            for c, h in enumerate(hdr, start=1):
                ws2.cell(row=5, column=c, value=h)
            # Blank days/hours in summary — must fall back to daily log totals
            data = ["Kiosk Test", badge, "Hourly", "10,000", "", "", 0, "", ""]
            for c, v in enumerate(data, start=1):
                ws2.cell(row=6, column=c, value=v)
            # Daily log sheet
            ws3 = wb2.create_sheet("Daily Log")
            d_hdr = ["Staff Name", "Badge Number", "Date", "Arrival", "Exit", "Hours", "Notes"]
            for c, h in enumerate(d_hdr, start=1):
                ws3.cell(row=3, column=c, value=h)
            # Two days with 9h each — precomputed to avoid formulas
            ws3.cell(row=4, column=1, value="Kiosk Test")
            ws3.cell(row=4, column=2, value=badge)
            ws3.cell(row=4, column=3, value="2026-02-01")
            ws3.cell(row=4, column=4, value="08:00")
            ws3.cell(row=4, column=5, value="17:00")
            ws3.cell(row=4, column=6, value=9)
            ws3.cell(row=5, column=3, value="2026-02-02")
            ws3.cell(row=5, column=4, value="08:00")
            ws3.cell(row=5, column=5, value="17:00")
            ws3.cell(row=5, column=6, value=9)
            buf = io.BytesIO(); wb2.save(buf); buf.seek(0)
            class FakeUploadFile:
                filename = "sheet.xlsx"
                async def read(self): return buf.getvalue()
            up_result = await timesheet_upload(
                file=FakeUploadFile(), period=now.strftime("%Y-%m"),
                location_id=loc_id, current_user=director,
            )
            assert up_result["created"], f"expected create, got {up_result}"
            assert up_result["created"][0]["hours_worked"] == 18.0
            latest = await db.hr_timesheets.find_one(
                {"staff_id": staff_id, "period": now.strftime("%Y-%m"), "source": "xlsx_upload"},
                {"_id": 0},
            )
            assert latest["hours_worked"] == 18.0
            assert "Daily Log" in (latest.get("notes") or "")
        finally:
            await db.users.delete_many({"id": {"$in": [staff_id, random_id]}})
            await db.hr_salaries.delete_one({"id": sal_id})
            await db.hr_timesheets.delete_many({"staff_id": staff_id})
            await db.hr_settings.delete_one({"location_id": loc_id})
    asyncio.run(_t())
